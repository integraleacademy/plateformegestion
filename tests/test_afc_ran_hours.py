from copy import deepcopy
from datetime import date
import json
from pathlib import Path
import sys

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as application
from services.afc_ran_hours import correct_afc_ran_hours

INTERRUPTIONS = [(date(2026, 12, 23), date(2027, 1, 4))]


def old_ran_planning():
    planning = application.build_afc_aps_ssiap_planning_data(date(2026, 11, 16), "Formateur", "Salle")
    afternoon = next(day for day in planning if day["date"] == "2026-11-24")["slots"][-1]
    afternoon.update(end="16:30", durationMinutes=180, duration=3)
    return planning


def test_ran_correction_removes_only_one_hour_and_preserves_all_other_activities():
    before = old_ran_planning()
    snapshot = deepcopy(before)
    after = correct_afc_ran_hours(before)
    assert before == snapshot
    changed = [(old, new) for old, new in zip(before, after) if old != new]
    assert len(changed) == 1
    old, new = changed[0]
    assert new["date"] == "2026-11-24"
    assert old["slots"][:-1] == new["slots"][:-1]
    assert new["slots"][-1]["end"] == "15:30"
    assert new["slots"][-1]["durationMinutes"] == 120
    summary = application.afc_aps_ssiap_summary_from_data(after, INTERRUPTIONS)
    assert summary["errors"] == []
    assert summary["uv_totals"]["RAN"] == 54
    assert summary["uv_totals"]["PAF"] == 20
    assert summary["total_hours"] == 392
    assert summary["days_count"] == 57
    assert correct_afc_ran_hours(after) == after
    assert application.AFC_DSF_MODULES["RAN"]["theoreticalHours"] == 54
    assert application.AFC_DSF_MODULES["PAF"]["theoreticalHours"] == 20


def test_old_volume_is_rejected_and_hour_cannot_be_transferred_to_paf():
    planning = old_ran_planning()
    errors = application.validate_actual_afc_hours(planning)
    assert any("55 h planifiées au lieu de 54 h" in error for error in errors)
    planning = correct_afc_ran_hours(planning)
    slot = next(slot for day in planning for slot in day["slots"] if slot["afcCategory"] == "PAF")
    slot["durationMinutes"] += 60
    errors = application.validate_actual_afc_hours(planning)
    assert any("21 h planifiées au lieu de 20 h" in error for error in errors)


def test_unexpected_ran_total_does_not_modify_input():
    planning = old_ran_planning()
    planning[0]["slots"][0]["durationMinutes"] -= 120
    before = deepcopy(planning)
    with pytest.raises(ValueError, match="ancien volume de 55h"):
        correct_afc_ran_hours(planning)
    assert planning == before


def stored_session(tmp_path, monkeypatch):
    planning_dir = tmp_path / "plannings"
    planning_dir.mkdir()
    monkeypatch.setattr(application, "DATA_DIR", str(tmp_path))
    monkeypatch.setattr(application, "SESSIONS_FILE", str(tmp_path / "sessions.json"))
    monkeypatch.setattr(application, "PLANNING_DIR", str(planning_dir))
    target = {"id": "9f823786", "formation": "AFC_APS_SSIAP", "training_code": "AFC_APS_SSIAP", "date_debut": "2026-11-16", "date_fin": "2027-02-15", "date_exam": "2027-02-15", "contractual_end_date": "2027-02-15", "interruptions": "23/12/2026 au 04/01/2027", "apsPlanningData": old_ran_planning(), "planning_pdf": "old.pdf", "apsAttendanceStudents": [{"id": "student"}], "afcDsfs": [{"id": "issued-dsf", "totalHours": 55}], "notes": "À conserver"}
    data = {"sessions": [target, {"id": "other", "formation": "APS"}], "jurys": [{"id": "jury"}]}
    application.save_sessions(data)
    (planning_dir / "old.pdf").write_bytes(b"original PDF")
    return data


def test_migration_pdf_and_exports_match_54_ran_20_paf_and_392_total(tmp_path, monkeypatch):
    from pypdf import PdfReader
    from openpyxl import load_workbook
    from services.afc_france_travail_attendance import generate_france_travail_workbook

    before = stored_session(tmp_path, monkeypatch)
    result = application.migrate_afc_ran_session()
    assert result == {"status": "applied", "session": "9f823786", "changed_dates": ["2026-11-24"], "ran_hours": 54, "paf_hours": 20, "hours": 392, "days": 57, "changed": True}
    after = application.load_sessions()
    target = after["sessions"][0]
    assert after["sessions"][1:] == before["sessions"][1:] and after["jurys"] == before["jurys"]
    for key in ("apsAttendanceStudents", "afcDsfs", "notes", "date_exam", "date_fin"):
        assert target[key] == before["sessions"][0][key]
    pdf = tmp_path / "plannings" / target["planning_pdf"]
    text = "\n".join(page.extract_text() for page in PdfReader(pdf).pages)
    assert "Parcours complet — 392 heures" in text
    assert "Remise à niveau (RAN) RAN 54 h" in " ".join(text.split())
    assert "Préparation à l’après-formation (PAF) PAF 20 h" in " ".join(text.split())
    assert "393" not in text and "55 h" not in text
    backup = next((tmp_path / "planning_migrations").rglob("*.json"))
    assert json.loads(backup.read_text()) == before["sessions"][0]
    assert backup.with_suffix(".pdf").read_bytes() == b"original PDF"
    snapshot = deepcopy(target)
    snapshot["apsAttendanceStudents"] = [{"id": "sample", "lastName": "TEST", "firstName": "Stagiaire"}]
    workbook = load_workbook(generate_france_travail_workbook(snapshot, Path(__file__).resolve().parents[1]))
    total = 0
    for sheet in workbook:
        row = next(row[0].row for row in sheet if row[0].value == "Total des heures facturables")
        total += float(str(sheet.cell(row, 13).value or 0).replace(",", "."))
    assert total == 392
    stored, generated = Path(application.SESSIONS_FILE).read_bytes(), pdf.read_bytes()
    assert application.migrate_afc_ran_session() == {"status": "already_applied"}
    assert Path(application.SESSIONS_FILE).read_bytes() == stored and pdf.read_bytes() == generated


def test_pdf_failure_preserves_previous_planning_and_pdf(tmp_path, monkeypatch):
    stored_session(tmp_path, monkeypatch)
    stored = Path(application.SESSIONS_FILE).read_bytes()
    def fail(*args, **kwargs):
        raise RuntimeError("PDF indisponible")
    monkeypatch.setattr(application, "generate_aps_planning_pdf", fail)
    with pytest.raises(RuntimeError, match="PDF indisponible"):
        application.migrate_afc_ran_session()
    assert Path(application.SESSIONS_FILE).read_bytes() == stored
    assert (tmp_path / "plannings" / "old.pdf").read_bytes() == b"original PDF"
    assert not list((tmp_path / "plannings").glob("afc_ran_*"))
