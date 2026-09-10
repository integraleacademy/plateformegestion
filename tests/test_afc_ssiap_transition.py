from collections import Counter
from copy import deepcopy
from datetime import date
import json
from pathlib import Path
import re
import sys

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as application
from services.afc_ssiap_transition import correct_saved_ssiap_transition

INTERRUPTIONS = [(date(2026, 12, 23), date(2027, 1, 4))]


def minutes(value):
    h, m = value.split(":")
    return int(h)*60 + int(m)


def legacy_planning():
    """Reconstituer l'ancienne répartition, avec 30 min de SSIAP le 20 janvier."""
    planning = application.build_afc_aps_ssiap_planning_data(date(2026, 11, 16), "Formateur", "Salle")
    windows = [("2027-01-20", 960, 990)]
    for day in planning:
        for slot in day["slots"]:
            if slot["afcCategory"] == "SSIAP1":
                start, end = minutes(slot["start"]), minutes(slot["end"])
                if day["date"] == "2027-02-03":
                    end = min(end, 930)
                if end > start:
                    windows.append((day["date"], start, end))
        day["slots"] = [slot for slot in day["slots"] if slot["afcCategory"] != "SSIAP1"]
    curriculum = [dict(module, remaining=module["durationMinutes"]) for module in application.afc_build_main_sequence() if module["category"] == "SSIAP1"]
    by_date = {day["date"]: day for day in planning}
    index = 0
    for iso, start, end in sorted(windows):
        while start < end:
            module = curriculum[index]
            duration = min(end-start, module["remaining"])
            by_date[iso]["slots"].append(application.afc_slot_from_module(date.fromisoformat(iso), start, start+duration, module, "Formateur", "Salle"))
            start += duration
            module["remaining"] -= duration
            if not module["remaining"]:
                index += 1
    for day in planning:
        day["slots"].sort(key=lambda slot: slot["start"])
    return planning


def curriculum_minutes(planning):
    totals = Counter()
    for day in planning:
        for slot in day["slots"]:
            if slot["afcCategory"] == "SSIAP1":
                content = {key: value for key, value in slot.items() if key not in {"start", "end", "duration", "durationMinutes"}}
                totals[json.dumps(content, sort_keys=True)] += slot["durationMinutes"]
    return totals


def test_saved_correction_preserves_programme_assignments_and_other_activities():
    original = legacy_planning()
    first = next(slot for day in original for slot in day["slots"] if slot["afcCategory"] == "SSIAP1")
    first.update(trainer="Intervenant feu", room="Salle pratique", notes="Consigne conservée")
    snapshot = deepcopy(original)
    corrected = correct_saved_ssiap_transition(original)
    assert original == snapshot
    assert curriculum_minutes(corrected) == curriculum_minutes(original)
    for before, after in zip(original, corrected):
        assert before["date"] == after["date"]
        assert [slot for slot in before["slots"] if slot["afcCategory"] != "SSIAP1"] == [slot for slot in after["slots"] if slot["afcCategory"] != "SSIAP1"]
    first_corrected = next((day["date"], slot) for day in corrected for slot in day["slots"] if slot["afcCategory"] == "SSIAP1")
    assert first_corrected[0] == "2027-01-22" and first_corrected[1]["start"] == "08:30"
    assert first_corrected[1]["notes"] == "Consigne conservée"
    assert application.afc_aps_ssiap_summary_from_data(corrected, INTERRUPTIONS)["errors"] == []
    assert correct_saved_ssiap_transition(corrected) == corrected


@pytest.mark.parametrize("bad_date", ["2027-01-20", "2027-01-21", "2027-01-26"])
def test_validation_rejects_ssiap_before_during_or_later_than_day_after_exam(bad_date):
    planning = application.build_afc_aps_ssiap_planning_data(date(2026, 11, 16))
    day22 = next(day for day in planning if day["date"] == "2027-01-22")
    slots = [slot for slot in day22["slots"] if slot["afcCategory"] == "SSIAP1"]
    day22["slots"] = [slot for slot in day22["slots"] if slot["afcCategory"] != "SSIAP1"]
    next(day for day in planning if day["date"] == bad_date)["slots"].extend(slots)
    errors = application.afc_aps_ssiap_summary_from_data(planning, INTERRUPTIONS)["errors"]
    assert any("Le SSIAP 1 doit commencer le lendemain" in error for error in errors)
    if bad_date <= "2027-01-21":
        assert any("avant ou pendant la journée d’examen APS" in error for error in errors)


def test_only_approved_dates_can_have_half_hour_billing():
    planning = application.build_afc_aps_ssiap_planning_data(date(2026, 11, 16))
    day18 = next(day for day in planning if day["date"] == "2027-01-18")
    day20 = next(day for day in planning if day["date"] == "2027-01-20")
    slot = day18["slots"].pop()
    assert slot["durationMinutes"] == 30
    slot.update(start="16:00", end="16:30")
    day20["slots"].append(slot)
    errors = application.afc_aps_ssiap_summary_from_data(planning, INTERRUPTIONS)["errors"]
    assert any("2027-01-18" in error and "heures entières" in error for error in errors)
    assert any("2027-01-20" in error and "doit durer 6h30" in error for error in errors)


def test_france_travail_export_preserves_both_half_hours_and_393_hour_total():
    from openpyxl import load_workbook
    from services.afc_france_travail_attendance import generate_france_travail_workbook

    session = {
        "id": "9f823786", "training_code": "AFC_APS_SSIAP",
        "date_debut": "2026-11-16", "date_fin": "2027-02-15",
        "apsPlanningData": application.build_afc_aps_ssiap_planning_data(date(2026, 11, 16)),
        "apsAttendanceStudents": [{"id": "one", "lastName": "TEST", "firstName": "Stagiaire"}],
    }
    workbook = load_workbook(generate_france_travail_workbook(session, Path(__file__).resolve().parents[1]))
    total_hours = 0
    for sheet in workbook:
        total_row = next(row[0].row for row in sheet if row[0].value == "Total des heures facturables")
        total_hours += float(str(sheet.cell(total_row, 13).value or 0).replace(",", "."))
        if sheet.title in {"1801 au 2201", "0102 au 0502"}:
            assert sheet["H12"].value == "13h30-16h00"
            assert sheet.cell(total_row, 7).value == 4
            assert sheet.cell(total_row, 8).value == "2,5"
    assert total_hours == 393


def saved_session(tmp_path, monkeypatch):
    planning_dir = tmp_path / "plannings"
    planning_dir.mkdir()
    monkeypatch.setattr(application, "DATA_DIR", str(tmp_path))
    monkeypatch.setattr(application, "SESSIONS_FILE", str(tmp_path / "sessions.json"))
    monkeypatch.setattr(application, "PLANNING_DIR", str(planning_dir))
    target = {"id": "9f823786", "formation": "AFC_APS_SSIAP", "training_code": "AFC_APS_SSIAP", "date_debut": "2026-11-16", "date_fin": "2027-02-15", "date_exam": "2027-02-15", "contractual_end_date": "2027-02-15", "interruptions": "23/12/2026 au 04/01/2027", "apsPlanningData": legacy_planning(), "planning_pdf": "old.pdf", "apsAttendanceStudents": [{"id": "student-to-preserve"}], "notes": "Conserver les autres champs"}
    data = {"sessions": [target, {"id": "other-session", "formation": "APS", "notes": "Ne pas modifier"}], "jurys": [{"id": "jury-to-preserve"}]}
    application.save_sessions(data)
    (planning_dir / "old.pdf").write_bytes(b"original PDF")
    return data


def test_persistent_migration_updates_only_target_with_backup_pdf_and_idempotence(tmp_path, monkeypatch):
    from pypdf import PdfReader

    before = saved_session(tmp_path, monkeypatch)
    result = application.migrate_afc_ssiap_transition_session()
    assert result == {"status": "applied", "session": "9f823786", "first_ssiap": ("2027-01-22", "08:30"), "hours": 393, "days": 57, "changed": True}
    after = application.load_sessions()
    target = after["sessions"][0]
    assert after["sessions"][1:] == before["sessions"][1:] and after["jurys"] == before["jurys"]
    assert target["apsAttendanceStudents"] == before["sessions"][0]["apsAttendanceStudents"]
    assert target["notes"] == before["sessions"][0]["notes"]
    assert not target["apsPlanningNeedsRegeneration"]
    pdf = tmp_path / "plannings" / target["planning_pdf"]
    reader = PdfReader(pdf)
    calendar = reader.pages[-1].extract_text()
    assert re.search(r"20\s+APS 4h\s+APS 2h30\s+21", calendar)
    assert re.search(r"3\s+SSIAP 1 4h\s+SSIAP 1 2h30\s+4", calendar)
    backup = next((tmp_path / "planning_migrations").rglob("*.json"))
    assert json.loads(backup.read_text()) == before["sessions"][0]
    assert backup.with_suffix(".pdf").read_bytes() == b"original PDF"
    stored_bytes = Path(application.SESSIONS_FILE).read_bytes()
    pdf_bytes = pdf.read_bytes()
    assert application.migrate_afc_ssiap_transition_session() == {"status": "already_applied"}
    assert Path(application.SESSIONS_FILE).read_bytes() == stored_bytes
    assert pdf.read_bytes() == pdf_bytes


def test_failed_pdf_generation_leaves_saved_session_and_existing_pdf_intact(tmp_path, monkeypatch):
    saved_session(tmp_path, monkeypatch)
    before = Path(application.SESSIONS_FILE).read_bytes()
    def fail(*args, **kwargs):
        raise RuntimeError("PDF indisponible")
    monkeypatch.setattr(application, "generate_aps_planning_pdf", fail)
    with pytest.raises(RuntimeError, match="PDF indisponible"):
        application.migrate_afc_ssiap_transition_session()
    assert Path(application.SESSIONS_FILE).read_bytes() == before
    assert (tmp_path / "plannings" / "old.pdf").read_bytes() == b"original PDF"
    assert not list((tmp_path / "plannings").glob("afc_transition_*"))
