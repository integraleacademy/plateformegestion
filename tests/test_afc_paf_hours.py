from collections import Counter
from copy import deepcopy
from datetime import date
import json
from pathlib import Path
import sys
from unittest.mock import patch

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as application
from services.afc_paf_hours import correct_afc_paf_hours
from services.afc_sst_planning import validate_afc_sst_days

INTERRUPTIONS = [(date(2026, 12, 23), date(2027, 1, 4))]


def old_paf_planning():
    # Reproduire le générateur précédent : quatre blocs de 5h, avant normalisation.
    with patch.object(application, "correct_afc_paf_hours", side_effect=lambda planning: planning), patch.object(application, "afc_aps_ssiap_summary_from_data", return_value={"errors": []}):
        return application.build_afc_aps_ssiap_planning_data(date(2026, 11, 16), "Formateur", "Salle")


def content_stream(planning):
    stream = []
    for day in planning:
        for slot in day["slots"]:
            if slot["afcCategory"] != "APS":
                continue
            content = {key: value for key, value in slot.items() if key not in {"start", "end", "duration", "durationMinutes"}}
            if stream and stream[-1][0] == content:
                stream[-1][1] += slot["durationMinutes"]
            else:
                stream.append([content, slot["durationMinutes"]])
    return stream


def test_paf_correction_preserves_dates_metadata_and_order_without_overloading_days():
    before = old_paf_planning()
    next(day for day in before if day["date"] == "2027-01-08")["slots"][0].update(trainer="Intervenant confirmé", notes="À conserver")
    snapshot = deepcopy(before)
    after = correct_afc_paf_hours(before)
    assert before == snapshot
    assert content_stream(after) == content_stream(before)
    for old, new in zip(before, after):
        assert old["date"] == new["date"]
        assert [s for s in old["slots"] if s["afcCategory"] not in {"APS", "PAF"}] == [s for s in new["slots"] if s["afcCategory"] not in {"APS", "PAF"}]
        if not "2027-01-06" <= old["date"] <= "2027-01-14":
            assert old == new
    paf = {day["date"]: sum(s["durationMinutes"] for s in day["slots"] if s["afcCategory"] == "PAF") for day in after if any(s["afcCategory"] == "PAF" for s in day["slots"])}
    assert paf == {"2027-01-14": 360, "2027-01-22": 300, "2027-01-28": 300, "2027-02-09": 300}
    assert validate_afc_sst_days(after) == []
    summary = application.afc_aps_ssiap_summary_from_data(after, INTERRUPTIONS)
    assert summary["errors"] == []
    assert summary["total_hours"] == 393
    assert summary["uv_totals"]["PAF"] == 21 and summary["uv_totals"]["RAN"] == 54
    assert Counter(sum(s["durationMinutes"] for s in d["slots"]) for d in after) == {420: 50, 360: 5, 390: 2}
    assert correct_afc_paf_hours(after) == after


def test_migrated_planning_matches_fresh_generation_and_rejects_old_volume():
    old = old_paf_planning()
    assert any("20 h planifiées au lieu de 21 h" in error for error in application.validate_actual_afc_hours(old))
    assert correct_afc_paf_hours(old) == application.build_afc_aps_ssiap_planning_data(date(2026, 11, 16), "Formateur", "Salle")
    assert application.AFC_DSF_MODULES["PAF"]["theoreticalHours"] == 21


@pytest.mark.parametrize("change", ["unexpected_volume", "occupied_hour", "atomic_aps"])
def test_unexpected_saved_schedule_is_preserved(change):
    planning = old_paf_planning()
    if change == "unexpected_volume":
        next(s for d in planning for s in d["slots"] if s["afcCategory"] == "PAF")["durationMinutes"] -= 60
    elif change == "occupied_hour":
        next(d for d in planning if d["date"] == "2027-01-06")["slots"][-1].update(end="16:30", durationMinutes=150)
    else:
        next(d for d in planning if d["date"] == "2027-01-08")["slots"][0]["is_atomic"] = True
    before = deepcopy(planning)
    with pytest.raises(ValueError):
        correct_afc_paf_hours(planning)
    assert planning == before


def stored_session(tmp_path, monkeypatch):
    planning_dir = tmp_path / "plannings"
    planning_dir.mkdir()
    monkeypatch.setattr(application, "DATA_DIR", str(tmp_path))
    monkeypatch.setattr(application, "SESSIONS_FILE", str(tmp_path / "sessions.json"))
    monkeypatch.setattr(application, "PLANNING_DIR", str(planning_dir))
    target = {"id": "9f823786", "formation": "AFC_APS_SSIAP", "training_code": "AFC_APS_SSIAP", "date_debut": "2026-11-16", "date_fin": "2027-02-15", "date_exam": "2027-02-15", "contractual_end_date": "2027-02-15", "interruptions": "23/12/2026 au 04/01/2027", "apsPlanningData": old_paf_planning(), "planning_pdf": "old.pdf", "apsAttendanceStudents": [{"id": "student"}], "afcDsfs": [{"id": "issued-dsf", "totalHours": 20}], "notes": "À conserver", "afcRanHoursVersion": "2026-10-26-ran54-paf20-v1"}
    data = {"sessions": [target, {"id": "other", "formation": "APS"}], "jurys": [{"id": "jury"}]}
    application.save_sessions(data)
    (planning_dir / "old.pdf").write_bytes(b"original PDF")
    return data


def test_migration_backs_up_saved_session_and_updates_pdf_and_exports(tmp_path, monkeypatch):
    from pypdf import PdfReader
    from openpyxl import load_workbook
    from services.afc_france_travail_attendance import generate_france_travail_workbook

    before = stored_session(tmp_path, monkeypatch)
    result = application.migrate_afc_paf_session()
    assert result == {"status": "applied", "session": "9f823786", "changed_dates": ["2027-01-06", "2027-01-08", "2027-01-11", "2027-01-14"], "ran_hours": 54, "paf_hours": 21, "hours": 393, "days": 57, "changed": True}
    after = application.load_sessions()
    target = after["sessions"][0]
    assert after["sessions"][1:] == before["sessions"][1:] and after["jurys"] == before["jurys"]
    for key in ("apsAttendanceStudents", "afcDsfs", "notes", "date_exam", "date_fin", "afcRanHoursVersion"):
        assert target[key] == before["sessions"][0][key]
    pdf = tmp_path / "plannings" / target["planning_pdf"]
    text = " ".join(" ".join(page.extract_text() for page in PdfReader(pdf).pages).split())
    assert "Parcours complet — 393 heures" in text
    assert "Remise à niveau (RAN) RAN 54 h" in text
    assert "Préparation à l’après-formation (PAF) PAF 21 h" in text
    backup = next((tmp_path / "planning_migrations").rglob("*.json"))
    assert json.loads(backup.read_text()) == before["sessions"][0]
    assert backup.with_suffix(".pdf").read_bytes() == b"original PDF"
    exported = deepcopy(target)
    exported["apsAttendanceStudents"] = [{"id": "sample", "lastName": "TEST", "firstName": "Stagiaire"}]
    workbook = load_workbook(generate_france_travail_workbook(exported, Path(__file__).resolve().parents[1]))
    total = 0
    for sheet in workbook:
        row = next(row[0].row for row in sheet if row[0].value == "Total des heures facturables")
        total += float(str(sheet.cell(row, 13).value or 0).replace(",", "."))
    assert total == 393
    stored, generated = Path(application.SESSIONS_FILE).read_bytes(), pdf.read_bytes()
    assert application.migrate_afc_paf_session() == {"status": "already_applied"}
    assert Path(application.SESSIONS_FILE).read_bytes() == stored and pdf.read_bytes() == generated


def test_failed_pdf_keeps_previous_session_and_document(tmp_path, monkeypatch):
    stored_session(tmp_path, monkeypatch)
    stored = Path(application.SESSIONS_FILE).read_bytes()
    def fail(*args, **kwargs):
        raise RuntimeError("PDF indisponible")
    monkeypatch.setattr(application, "generate_aps_planning_pdf", fail)
    with pytest.raises(RuntimeError, match="PDF indisponible"):
        application.migrate_afc_paf_session()
    assert Path(application.SESSIONS_FILE).read_bytes() == stored
    assert (tmp_path / "plannings" / "old.pdf").read_bytes() == b"original PDF"
    assert not list((tmp_path / "plannings").glob("afc_paf_*"))
